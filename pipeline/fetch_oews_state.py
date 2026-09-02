"""
BLS OEWS (Occupational Employment and Wage Statistics) state + national time series, 2018-2024.

Sources (one zip per year, each containing a single .xlsx data sheet):
    https://www.bls.gov/oes/special-requests/oesm{YY}st.zip    state files  (~6-8 MB)
    https://www.bls.gov/oes/special-requests/oesm{YY}nat.zip   national     (~0.3 MB)

Outputs:
    data/oews_state.json        54 areas (50 states + DC + PR/GU/VI) x all occupations x 7 years
    data/oews_national.json     area 99 "U.S." x all occupations x 7 years
    data/oews_state_major.json  ONLY if oews_state.json exceeds SIZE_LIMIT_MB

Everything is columnar: an `areas` index, an `occupations` index, and per-year parallel arrays of
row values. Nothing is derived that a browser could compute itself; jobs_1000 and loc_quotient are
carried as published because they are not recomputable from this file alone.

Upstream quirks handled explicitly (all verified against the real files, not assumed):
  * Header casing/naming drifts:   2018 = UPPER + legacy names (ST/STATE/OCC_GROUP/LOC_Q),
                                   2019 = lower, 2020-2024 = UPPER (AREA_TITLE/O_GROUP/LOC_QUOTIENT).
  * The 2018 NATIONAL file has no AREA/AREA_TITLE columns at all - synthesized as 99 / "U.S.".
  * The data sheet is located by looking for an `occ_code` header, never by sheet index or name
    (names drift: state_dl / State_M2020_dl / "All May 2021 data" / state_M2024_dl).
  * oesm22st.zip contains a stray Excel lock file `~$state_M2022_dl.xlsx` - skipped.
  * The NATIONAL files publish 7-13 occ_codes TWICE per year: once as o_group "broad" and once as
    "detailed", with identical values (a broad group whose only detailed occupation shares its
    code, e.g. 13-1020 Buyers and Purchasing Agents).  occ_code alone is therefore NOT a key.
    The occupation index is keyed by (occ_code, o_group); meta.duplicate_codes lists them so a
    consumer summing detailed rows does not double-count.  State files have no such duplicates.
  * BLS non-numeric markers map to null and are recorded in `flags`, never coerced to 0:
        **  estimate not released (employment / jobs_1000 / loc_quotient)
        *   wage estimate not available
        #   wage AT OR ABOVE the top-code (>= $115.00/hr or $239,200/yr in recent years)
        ~   (handled, not observed in these files)
  * OCCUPATION CODES ARE NOT COMPARABLE ACROSS THE 2018/2019 BOUNDARY, and some codes churn again
    at 2020/2021. OEWS moved from SOC 2010 (May 2018) to SOC 2018 (May 2019 onward), and several
    SOC 2018 codes were published only in combined form for 2019-2020. The software-developer
    lineage, verified in these files, is:
        2018        15-1132 Software Developers, Applications        (903,160 US)
                  + 15-1133 Software Developers, Systems Software    (405,330 US)
        2019-2020   15-1256 Software Developers and Software Quality Assurance Analysts and Testers
        2021-2024   15-1252 Software Developers  +  15-1253 Software QA Analysts and Testers
    Same story for web developers: 15-1257 (2019-2020) -> 15-1254 + 15-1255 (2021+).
    Every vintage is kept; meta.soc_vintage says which SOC edition each year uses, and
    occupations[].years says exactly which years each code was actually published in.
  * May 2020 reflects a COVID-affected reference period - flagged in meta.covid_years.
"""

import io
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pipeline.common import DATA, FIPS_TO_ABBR, RAW, YEARS, fetch, write_json  # noqa: E402

RAW_OEWS = RAW / "oews"
SIZE_LIMIT_MB = 25.0

# Retry header set for BLS, which 403s some clients.
ALT_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept": "application/zip,*/*",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": "https://www.bls.gov/oes/tables.htm",
}

# FIPS -> postal for territories OEWS publishes that common.STATES does not carry.
TERRITORY_FIPS = {"60": "AS", "66": "GU", "69": "MP", "72": "PR", "78": "VI"}

# Every value we keep, and the precision BLS publishes it at (None = integer).
NUM_FIELDS = {
    "tot_emp": None,
    "jobs_1000": 3,
    "loc_quotient": 2,
    "a_median": None,
    "a_mean": None,
    "h_median": 2,
}

# Header aliases, lower-cased. First match wins.
ALIASES = {
    "area": ("area",),
    "area_title": ("area_title", "state"),
    "area_type": ("area_type",),
    "occ_code": ("occ_code",),
    "occ_title": ("occ_title",),
    "o_group": ("o_group", "occ_group"),
    "tot_emp": ("tot_emp",),
    "jobs_1000": ("jobs_1000",),
    "loc_quotient": ("loc_quotient", "loc_q"),
    "a_median": ("a_median",),
    "a_mean": ("a_mean",),
    "h_median": ("h_median",),
}

SUPPRESSION = {"*", "**", "#", "~"}


# --------------------------------------------------------------------------------------- download

def download(year, kind):
    """kind is 'st' or 'nat'. Cached on disk; BLS is retried with browser-ish headers."""
    yy = f"{year % 100:02d}"
    url = f"https://www.bls.gov/oes/special-requests/oesm{yy}{kind}.zip"
    dest = RAW_OEWS / f"oesm{yy}{kind}.zip"
    try:
        return fetch(url, dest)
    except Exception as first:  # noqa: BLE001 - BLS blocks some UAs; try a second header set
        print(f"  [warn] {dest.name}: {first!r} -> retrying with alternate headers")
        return fetch(url, dest, headers=ALT_HEADERS)


# ------------------------------------------------------------------------------------------ parse

def open_data_sheet(zip_path):
    """Return (workbook, worksheet, lowercased header list, member name) for the occ_code sheet."""
    with zipfile.ZipFile(zip_path) as zf:
        members = [
            n for n in zf.namelist()
            if n.lower().endswith((".xlsx", ".xls"))
            and "field_desc" not in n.lower()
            and not Path(n).name.startswith("~$")   # stray Excel lock file in oesm22st.zip
        ]
        if not members:
            raise RuntimeError(f"{zip_path.name}: no data workbook inside")
        if len(members) > 1:
            print(f"  [note] {zip_path.name} has several workbooks: {members}")
        for member in members:
            wb = openpyxl.load_workbook(io.BytesIO(zf.read(member)), read_only=True, data_only=True)
            for ws in wb.worksheets:
                first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
                if not first:
                    continue
                hdr = [str(c).strip().lower() if c is not None else "" for c in first]
                if "occ_code" in hdr:
                    return wb, ws, hdr, member
            wb.close()
    raise RuntimeError(f"{zip_path.name}: no sheet with an occ_code column")


def column_index(hdr, logical):
    for alias in ALIASES[logical]:
        if alias in hdr:
            return hdr.index(alias)
    return None


def parse_number(raw, ndigits):
    """-> (value, marker). marker is the BLS suppression / top-code symbol, or None."""
    if raw is None:
        return None, None
    if isinstance(raw, str):
        s = raw.strip()
        if s == "":
            return None, None
        if s in SUPPRESSION:
            return None, s
        s = s.replace(",", "").replace("$", "")
        try:
            raw = float(s)
        except ValueError:
            return None, s  # unknown marker - preserved rather than silently dropped
    if ndigits is None:
        return int(round(float(raw))), None
    return round(float(raw), ndigits), None


def read_year(zip_path, *, national):
    """Return (normalized row dicts, workbook member name) for one OEWS year."""
    wb, ws, hdr, member = open_data_sheet(zip_path)
    idx = {k: column_index(hdr, k) for k in ALIASES}
    missing = [k for k in ("occ_code", "occ_title", "o_group") if idx[k] is None]
    if missing:
        raise RuntimeError(f"{zip_path.name}/{member}: missing required columns {missing} in {hdr}")

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r is None or idx["occ_code"] >= len(r) or r[idx["occ_code"]] is None:
            continue
        occ_code = str(r[idx["occ_code"]]).strip()
        if not occ_code:
            continue

        if idx["area"] is not None:
            area = str(r[idx["area"]]).strip()
            if area.isdigit() and len(area) < 2:   # xlsx may store FIPS as a number: 1 -> "01"
                area = area.zfill(2)
        else:
            area = "99"                            # 2018 national file has no AREA column
        area_title = str(r[idx["area_title"]]).strip() if idx["area_title"] is not None else "U.S."
        area_type = str(r[idx["area_type"]]).strip() if idx["area_type"] is not None else None

        rec = {
            "area": area,
            "area_title": area_title,
            "area_type": area_type,
            "occ_code": occ_code,
            "occ_title": str(r[idx["occ_title"]]).strip(),
            "o_group": str(r[idx["o_group"]]).strip().lower(),
        }
        for field, ndigits in NUM_FIELDS.items():
            col = idx[field]
            rec[field] = (None, None) if col is None else parse_number(r[col], ndigits)
        rows.append(rec)
    wb.close()

    if national:
        bad = {r["area"] for r in rows} - {"99"}
        if bad:
            raise RuntimeError(f"{zip_path.name}: national file has unexpected areas {sorted(bad)}")
    return rows, member


# ------------------------------------------------------------------------------------------ build

def area_kind(fips, area_type):
    """'state' (DC included, per OEWS AREA_TYPE 2), 'territory', or 'nation'."""
    if fips == "99":
        return "nation"
    if area_type == "3":
        return "territory"
    if area_type == "2":
        return "state"
    return "territory" if fips in TERRITORY_FIPS else "state"   # 2018 has no AREA_TYPE column


def build(per_year, *, national):
    """per_year: {year: [rec, ...]} -> (columnar object, areas, occupations)."""
    # ---- area index (stable order: sorted by FIPS) --------------------------------------------
    area_meta = {}
    for year in sorted(per_year):
        for rec in per_year[year]:
            # later years overwrite, so the newest published naming wins
            area_meta[rec["area"]] = (rec["area_title"], rec["area_type"])
    areas = []
    for fips in sorted(area_meta):
        title, atype = area_meta[fips]
        entry = {"fips": fips, "title": title, "kind": area_kind(fips, atype)}
        abbr = FIPS_TO_ABBR.get(fips) or TERRITORY_FIPS.get(fips)
        if abbr:
            entry["abbr"] = abbr
        areas.append(entry)
    area_pos = {a["fips"]: i for i, a in enumerate(areas)}

    # ---- occupation index ---------------------------------------------------------------------
    # Keyed by (occ_code, o_group): the national files publish some codes twice, once as "broad"
    # and once as "detailed" with identical values, so occ_code alone is not a key.
    titles_by_year = defaultdict(dict)      # (code, group) -> {year: title}
    groups_of_code = defaultdict(lambda: defaultdict(list))   # code -> group -> [years]
    for year in sorted(per_year):
        seen = set()
        for rec in per_year[year]:
            key = (rec["occ_code"], rec["o_group"])
            dup_key = (rec["area"], *key)
            if dup_key in seen:
                raise RuntimeError(
                    f"{year}: repeated (area, occ_code, o_group) {dup_key} - the chosen key is "
                    f"not unique upstream; the parser must be revisited before trusting output"
                )
            seen.add(dup_key)
            titles_by_year[key][year] = rec["occ_title"]
            if year not in groups_of_code[rec["occ_code"]][rec["o_group"]]:
                groups_of_code[rec["occ_code"]][rec["o_group"]].append(year)

    occupations, title_history = [], {}
    for key in sorted(titles_by_year):
        code, group = key
        by_year = titles_by_year[key]
        newest = max(by_year)
        occupations.append({
            "code": code,
            "title": by_year[newest],     # most recent year's published title
            "group": group,
            "years": sorted(by_year),
        })
        if len(set(by_year.values())) > 1:
            title_history.setdefault(code, {}).update(
                {str(y): t for y, t in sorted(by_year.items())}
            )
    occ_pos = {(o["code"], o["group"]): i for i, o in enumerate(occupations)}

    # Codes published under more than one o_group - real, and a double-counting trap.
    duplicate_codes = {
        code: {g: yrs for g, yrs in sorted(groups.items())}
        for code, groups in sorted(groups_of_code.items()) if len(groups) > 1
    }

    # ---- per-year value columns ---------------------------------------------------------------
    values = {}
    for year in sorted(per_year):
        recs = per_year[year]
        col = {"o": [occ_pos[(r["occ_code"], r["o_group"])] for r in recs]}
        if not national:
            col["s"] = [area_pos[r["area"]] for r in recs]
        flags = {}
        for field in NUM_FIELDS:
            vals = [r[field][0] for r in recs]
            marks = defaultdict(list)
            for i, r in enumerate(recs):
                m = r[field][1]
                if m is not None:
                    marks[m].append(i)
            if national and not marks and all(v is None for v in vals):
                continue  # national files carry no jobs_1000 / loc_quotient at all
            col[field] = vals
            if marks:
                flags[field] = {m: rows for m, rows in sorted(marks.items())}
        if flags:
            col["flags"] = flags
        values[str(year)] = col

    out = {
        "meta": {
            "source": "BLS OEWS (Occupational Employment and Wage Statistics), May of each year",
            "level": "national" if national else "state",
            "urls": {
                str(y): "https://www.bls.gov/oes/special-requests/"
                        f"oesm{y % 100:02d}{'nat' if national else 'st'}.zip"
                for y in sorted(per_year)
            },
            "row_order": "as published upstream; join through the o/s index arrays, not row order",
            "null_means": "value not published upstream; see flags for the reason",
            "flag_legend": {
                "**": "estimate not released",
                "*": "wage estimate not available",
                "#": "wage at or above the BLS top-code (>= $115.00/hr or $239,200/yr recently)",
                "~": "value less than 0.005 percent",
            },
            "precision": {
                k: ("integer" if v is None else f"{v} decimals") for k, v in NUM_FIELDS.items()
            },
            "covid_years": [2020],
            "covid_note": "May 2020 OEWS reflects a COVID-affected reference period.",
            "soc_vintage": {"2018": "SOC 2010", **{str(y): "SOC 2018" for y in range(2019, 2025)}},
            "soc_note": (
                "Occupation codes are NOT comparable across the 2018/2019 boundary. OEWS switched "
                "from SOC 2010 (May 2018) to SOC 2018 (May 2019+). e.g. 15-1132 + 15-1133 (2018) "
                "-> 15-1252 Software Developers (2021+). Use occupations[].years to see exactly "
                "which years each code was published in; codes do come and go mid-series."
            ),
            "occupation_key": (
                "occupations[] is keyed by (code, group), not code alone: the national files "
                "publish some codes twice, once as broad and once as detailed with identical "
                "values. See duplicate_codes."
            ),
        },
        "years": sorted(per_year),
        "occupations": occupations,
        "values": values,
    }
    if not national:
        out["areas"] = areas
    if title_history:
        out["title_history"] = title_history
    if duplicate_codes:
        out["duplicate_codes"] = duplicate_codes
    return out, areas, occupations


# ----------------------------------------------------------------------------------------- sanity

def sanity(per_year, *, national, label):
    print(f"\n--- sanity: {label} ---")
    for year in sorted(per_year):
        recs = per_year[year]
        tots = [r for r in recs if r["occ_code"] == "00-0000"]
        if national:
            emp = tots[0]["tot_emp"][0] if tots else None
            print(f"  {year}  US 00-0000 tot_emp = {emp:,}" if emp is not None
                  else f"  {year}  US total MISSING")
        else:
            st = [r for r in tots if area_kind(r["area"], r["area_type"]) == "state"]
            te = [r for r in tots if area_kind(r["area"], r["area_type"]) == "territory"]
            s = sum(r["tot_emp"][0] for r in st if r["tot_emp"][0] is not None)
            top = max(st, key=lambda r: r["tot_emp"][0] or -1)
            tsum = sum(r["tot_emp"][0] for r in te if r["tot_emp"][0] is not None)
            print(f"  {year}  sum of {len(st)} states+DC 00-0000 = {s:,}"
                  f"   largest = {top['area_title']} {top['tot_emp'][0]:,}"
                  f"   ({len(te)} territories, {tsum:,})")

    dev_codes = {
        "15-1252": "Software Developers (SOC 2018)",
        "15-1132": "Software Developers, Applications (SOC 2010)",
        "15-1133": "Software Developers, Systems Software (SOC 2010)",
    }
    scope_area, scope = ("99", "US") if national else ("06", "CA")
    for code, name in dev_codes.items():
        line = []
        for year in sorted(per_year):
            hits = [r for r in per_year[year]
                    if r["occ_code"] == code and r["area"] == scope_area]
            if hits:
                v = hits[0]["tot_emp"][0]
                line.append(f"{year}:{v:,}" if v is not None else f"{year}:null")
        if line:
            print(f"  {code} {name} [{scope}]  " + "  ".join(line))


# ------------------------------------------------------------------------------------------- main

def emit_major_subset(obj):
    """Reduced state file: total + major groups + every detailed 15-xxxx occupation."""
    old = obj["occupations"]
    keep = {i for i, o in enumerate(old)
            if o["group"] in ("total", "major") or o["code"].startswith("15-")}
    dropped = [o["code"] for i, o in enumerate(old) if i not in keep]

    sub = {k: v for k, v in obj.items() if k not in ("occupations", "values", "meta")}
    sub["meta"] = dict(obj["meta"])
    sub["meta"]["subset"] = (
        "RESTRICTED FILE: total (00-0000) + all major groups (XX-0000) + every detailed 15-xxxx "
        f"(Computer & Mathematical) occupation. {len(dropped)} other detailed occupations are "
        "excluded here; the complete matrix is data/oews_state.json."
    )
    sub["occupations"] = [o for i, o in enumerate(old) if i in keep]
    pos = {oi: j for j, oi in enumerate(sorted(keep))}

    sub["values"] = {}
    for y, col in obj["values"].items():
        take = [i for i, oi in enumerate(col["o"]) if oi in keep]
        remap = {i: j for j, i in enumerate(take)}
        nc = {"o": [pos[col["o"][i]] for i in take],
              "s": [col["s"][i] for i in take]}
        for f in NUM_FIELDS:
            if f in col:
                nc[f] = [col[f][i] for i in take]
        if "flags" in col:
            nf = {}
            for f, marks in col["flags"].items():
                mm = {m: [remap[i] for i in rows if i in remap] for m, rows in marks.items()}
                mm = {m: v for m, v in mm.items() if v}
                if mm:
                    nf[f] = mm
            if nf:
                nc["flags"] = nf
        sub["values"][y] = nc
    return sub, dropped


def main():
    RAW_OEWS.mkdir(parents=True, exist_ok=True)
    report, failed = {}, []

    for kind, national, out_name in (("st", False, "oews_state.json"),
                                     ("nat", True, "oews_national.json")):
        label = "national" if national else "state"
        per_year, rows_in = {}, 0
        print(f"\n=== OEWS {label} ===")
        for year in YEARS:
            try:
                zp = download(year, kind)
            except Exception as e:  # noqa: BLE001
                print(f"  {year}: DOWNLOAD FAILED {e!r}")
                failed.append((label, year, f"download: {e!r}"))
                continue
            try:
                recs, member = read_year(zp, national=national)
            except Exception as e:  # noqa: BLE001
                print(f"  {year}: PARSE FAILED {e!r}")
                failed.append((label, year, f"parse: {e!r}"))
                continue
            per_year[year] = recs
            rows_in += len(recs)
            groups = Counter(r["o_group"] for r in recs)
            print(f"  {year}: {zp.name}/{Path(member).name}  rows={len(recs):,}  "
                  f"areas={len({r['area'] for r in recs})}  groups={dict(sorted(groups.items()))}")

        if not per_year:
            print(f"  no {label} years parsed - no output written")
            continue

        obj, areas, occs = build(per_year, national=national)
        sanity(per_year, national=national, label=label)

        rows_out = sum(len(v["o"]) for v in obj["values"].values())
        path = write_json(obj, DATA / out_name)
        mb = path.stat().st_size / 1024 / 1024
        print(f"  rows in={rows_in:,}  rows out={rows_out:,}  occupations={len(occs)}  "
              f"areas={len(areas)}  size={mb:.2f} MB")
        if obj.get("title_history"):
            print(f"  occ codes whose title changed across years: {len(obj['title_history'])}")
        if obj.get("duplicate_codes"):
            dc = obj["duplicate_codes"]
            print(f"  occ codes published under >1 o_group (kept as separate index entries, "
                  f"do NOT sum both): {len(dc)} -> {list(dc)[:6]}")
        report[label] = {"path": path, "mb": mb, "rows_in": rows_in, "rows_out": rows_out,
                         "occs": len(occs), "areas": len(areas), "obj": obj}

    # ---- optional reduced state file -----------------------------------------------------------
    st = report.get("state")
    if st and st["mb"] > SIZE_LIMIT_MB:
        print(f"\n=== oews_state.json is {st['mb']:.2f} MB > {SIZE_LIMIT_MB} MB: also emitting "
              f"data/oews_state_major.json ===")
        sub, dropped = emit_major_subset(st["obj"])
        p2 = write_json(sub, DATA / "oews_state_major.json")
        print(f"  EXCLUDED from the reduced file only: {len(dropped)} detailed non-15-xxxx "
              f"occupations (e.g. {dropped[:5]})")
        print(f"  reduced rows out={sum(len(v['o']) for v in sub['values'].values()):,}  "
              f"size={p2.stat().st_size / 1024 / 1024:.2f} MB")
    elif st:
        print(f"\n=== oews_state.json is {st['mb']:.2f} MB <= {SIZE_LIMIT_MB} MB: reduced file not "
              f"needed; NOTHING was excluded from any output ===")

    # ---- summary -------------------------------------------------------------------------------
    print("\n=== SUMMARY ===")
    for label in ("state", "national"):
        r = report.get(label)
        if not r:
            print(f"  {label}: NO OUTPUT")
            continue
        yrs = sorted(int(y) for y in r["obj"]["values"])
        print(f"  {label}: {r['path'].name}  years={yrs}  rows_in={r['rows_in']:,}  "
              f"rows_out={r['rows_out']:,}  occupations={r['occs']}  areas={r['areas']}  "
              f"{r['mb']:.2f} MB")
        for y in yrs:
            c = r["obj"]["values"][str(y)]
            nn = sum(1 for v in c.get("tot_emp", []) if v is not None)
            print(f"      {y}: {len(c['o']):,} rows; tot_emp non-null {nn:,}, "
                  f"not released {len(c['o']) - nn:,}")
    if failed:
        print("  YEARS THAT FAILED:")
        for label, year, why in failed:
            print(f"    {label} {year}: {why}")
    else:
        print("  all requested years (2018-2024, state + national) fetched and parsed")


if __name__ == "__main__":
    main()
